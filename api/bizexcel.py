"""
전체 사업 현황 엑셀 생성 (/api/bizexcel)

accounting.html의 '전체 사업 현황 > 엑셀 저장' 버튼이 호출한다.
브라우저(SheetJS)로는 굵은 글씨·테두리·색 같은 서식을 못 넣어서,
서버에서 openpyxl로 기존 '2026년도 사업별 사업 업무진행 현황' 양식 그대로 만든다.

요청 (POST, JSON):
  {
    "asOf": "2026.09.22",
    "rows": [{
       "type": "국고보조금", "manager": "우형", "agency": "법무부(보장원)",
       "target": 250, "apply": 102, "done": 25,
       "budget": 500000000, "paid": 400000000, "exec": 261804990,
       "note": "1분기 교부신청 제출\n...",
       "cats": [{"cat":"사업비","sub":"소송구조","budget":287800000}, ...]
    }, ...]
  }
응답: { "xlsx": "<base64>" }
"""

import io, json, base64
from http.server import BaseHTTPRequestHandler

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GREEN = "0F6E56"
LIGHT = "EEF3F1"
HEADB = "DCE6E2"
LINE  = "BFCCC7"
WARN  = "B45309"

THIN  = Side(style="thin",   color=LINE)
MED   = Side(style="medium", color=GREEN)
BOX   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FN = "맑은 고딕"
F_TITLE = Font(name=FN, size=14, bold=True, color="1F2937")
F_ASOF  = Font(name=FN, size=9,  color="6B7280")
F_HEAD  = Font(name=FN, size=9,  bold=True, color="FFFFFF")
F_BODY  = Font(name=FN, size=9)
F_BOLD  = Font(name=FN, size=9,  bold=True)
F_SMALL = Font(name=FN, size=8,  color="6B7280")
F_WARN  = Font(name=FN, size=9,  color=WARN)

FILL_HEAD = PatternFill("solid", fgColor=GREEN)
FILL_SUM  = PatternFill("solid", fgColor=HEADB)

A_CW = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_C  = Alignment(horizontal="center", vertical="center")
A_LW = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
A_L  = Alignment(horizontal="left",   vertical="center")
A_R  = Alignment(horizontal="right",  vertical="center")

NUM = '#,##0'
PCT = '0.0%'

# 열 구성 — 기존 양식과 같은 순서
HEADS = ["사업\n유형", "담당자", "지원부처", "사업 세부예산내역",
         "사업비", "운영비", "일반관리비", "목표\n건수",
         "예산", "교부액", "미교부", "신청\n건수", "집행\n건수",
         "현 집행액", "현 잔액", "예산 총 잔액", "사업 진행내역"]
WIDTHS = [7, 8, 17, 30, 14, 13, 13, 7, 15, 15, 14, 7, 7, 15, 14, 15, 34]
MONEY_COLS = [5, 6, 7, 9, 10, 11, 14, 15, 16]   # 천단위 서식 적용 열


def _n(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _i(v):
    try:
        return int(v or 0)
    except (TypeError, ValueError):
        return 0


def _group(cats):
    """비목 이름으로 사업비 / 운영비 / 일반관리비 소계를 만든다."""
    s = {"사업비": 0.0, "운영비": 0.0, "일반관리비": 0.0}
    for c in cats or []:
        name = str(c.get("cat") or "")
        amt  = _n(c.get("budget"))
        if "일반관리" in name:
            s["일반관리비"] += amt
        elif "운영" in name or "인건" in name:
            s["운영비"] += amt
        else:
            s["사업비"] += amt
    return s


def _detail(cats):
    """세부예산내역 글자. 항목 설정(BUDGET_CATS)에서 그대로 뽑으므로 항상 최신."""
    out = []
    for c in cats or []:
        sub = str(c.get("sub") or c.get("cat") or "").strip()
        if not sub:
            continue
        out.append("- {} {:,}".format(sub, int(_n(c.get("budget")))))
    return "\n".join(out)


def build(payload):
    as_of = str(payload.get("asOf") or "")
    rows  = payload.get("rows") or []
    ncol  = len(HEADS)
    last  = get_column_letter(ncol)

    wb = Workbook()
    ws = wb.active
    ws.title = "사업현황"
    ws.sheet_view.showGridLines = False

    # ── 제목 ──
    ws.merge_cells(f"A1:{last}1")
    t = ws.cell(row=1, column=1, value="2026년도 사업별 사업 업무진행 현황")
    t.font, t.alignment = F_TITLE, A_C
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{last}2")
    a = ws.cell(row=2, column=1, value=("기준일: " + as_of) if as_of else "")
    a.font, a.alignment = F_ASOF, A_R

    # ── 머리글 ──
    for i, (label, w) in enumerate(zip(HEADS, WIDTHS), start=1):
        c = ws.cell(row=3, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, A_CW, BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 34

    # ── 본문: 사업마다 값 행 + 비율 행 2줄 ──
    r = 4
    T = {"budget": 0.0, "paid": 0.0, "exec": 0.0,
         "사업비": 0.0, "운영비": 0.0, "일반관리비": 0.0,
         "target": 0, "apply": 0, "done": 0}

    for b in rows:
        cats   = b.get("cats") or []
        g      = _group(cats)
        budget = _n(b.get("budget")); paid = _n(b.get("paid")); ex = _n(b.get("exec"))
        target = _i(b.get("target")); done = _i(b.get("done")); apply_ = _i(b.get("apply"))

        vals = [
            b.get("type") or "", b.get("manager") or "", b.get("agency") or "",
            _detail(cats),
            g["사업비"] or None, g["운영비"] or None, g["일반관리비"] or None,
            target or None,
            budget, paid, budget - paid,
            apply_ or None, done or None,
            ex, paid - ex, budget - ex,
            b.get("note") or "",
        ]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border, c.font = BOX, F_BODY
            if i in (1, 2, 3):
                c.alignment = A_CW
            elif i == 4 or i == ncol:
                c.alignment = A_LW
            elif i in MONEY_COLS:
                c.alignment, c.number_format = A_R, NUM
                if isinstance(v, (int, float)) and v < 0:
                    c.font = F_WARN
            else:
                c.alignment = A_C

        # 비율 행 (달성률 / 집행률 / 잔액비율)
        r2 = r + 1
        ratios = {
            13: (done / target) if target else None,          # 집행건수 달성률
            14: (ex / budget) if budget else None,             # 집행률
            15: ((paid - ex) / budget) if budget else None,    # 현 잔액 비율
            16: ((budget - ex) / budget) if budget else None,  # 총 잔액 비율
        }
        for i in range(1, ncol + 1):
            c = ws.cell(row=r2, column=i, value=ratios.get(i))
            c.border, c.font, c.alignment = BOX, F_SMALL, A_R
            if i in ratios and ratios[i] is not None:
                c.number_format = PCT
            else:
                c.alignment = A_C

        # 값 행 + 비율 행을 하나로 보이게 앞쪽 칸은 세로 병합
        for col in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, ncol):
            ws.merge_cells(start_row=r, start_column=col, end_row=r2, end_column=col)

        ws.row_dimensions[r].height = 62
        ws.row_dimensions[r2].height = 14

        T["budget"] += budget; T["paid"] += paid; T["exec"] += ex
        T["사업비"] += g["사업비"]; T["운영비"] += g["운영비"]; T["일반관리비"] += g["일반관리비"]
        T["target"] += target; T["apply"] += apply_; T["done"] += done
        r += 2

    # ── 합계 ──
    tot = ["합 계", "", "", "",
           T["사업비"] or None, T["운영비"] or None, T["일반관리비"] or None,
           T["target"] or None,
           T["budget"], T["paid"], T["budget"] - T["paid"],
           T["apply"] or None, T["done"] or None,
           T["exec"], T["paid"] - T["exec"], T["budget"] - T["exec"], ""]
    for i, v in enumerate(tot, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.border, c.font, c.fill = BOX, F_BOLD, FILL_SUM
        if i in MONEY_COLS:
            c.alignment, c.number_format = A_R, NUM
        elif i == 1:
            c.alignment = A_C
        else:
            c.alignment = A_C
        c.border = Border(left=THIN, right=THIN, top=MED, bottom=THIN)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 22

    ws.freeze_panes = "D4"
    ws.print_title_rows = "3:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


class handler(BaseHTTPRequestHandler):
    def _send(self, code, obj):
        body = json.dumps(obj, ensure_ascii=False).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        try:
            n = int(self.headers.get("Content-Length") or 0)
            payload = json.loads(self.rfile.read(n) or b"{}")
            self._send(200, {"xlsx": base64.b64encode(build(payload)).decode()})
        except Exception as e:
            import traceback
            self._send(500, {"error": str(e), "trace": traceback.format_exc()})

    def do_GET(self):
        self._send(200, {"ok": True, "info": "POST JSON to build the business-status workbook"})
